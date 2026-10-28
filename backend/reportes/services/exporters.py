import io
import csv
from datetime import datetime
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelReportBuilder:
    """
    Generador de archivos Excel estilizados bajo los estándares visuales de Sistema TM:
    Azul corporativo (#1E40AF), acentos dorados (#D97706), bordes limpios y autoajuste.
    """

    COLOR_PRIMARY = "1E40AF"     # Azul TM
    COLOR_ACCENT = "D97706"      # Ámbar / Amarillo TM
    COLOR_BG_HEADER = "1E40AF"
    COLOR_TEXT_HEADER = "FFFFFF"
    COLOR_ROW_ALT = "F8FAFC"     # Gris/Azul tenue alternado
    COLOR_TOTAL = "E2E8F0"       # Fila de totales

    @staticmethod
    def build_excel(title, headers, rows, summary=None, metadata=None):
        """
        title: str (ej. 'Reporte de Stock Crítico')
        headers: list of str (ej. ['Código', 'Nombre', 'Stock', 'Stock Mínimo', 'Estado'])
        rows: list of lists (valores de cada fila)
        summary: dict or list of dicts (ej. [{'label': 'Total Repuestos', 'value': 45}])
        metadata: dict (ej. {'Usuario': 'admin', 'Rango Fechas': '2026-01-01 a 2026-10-26'})
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:30]  # Limite de longitud en nombres de hoja Excel

        # Estilos base
        font_title = Font(name="Calibri", size=15, bold=True, color="1E40AF")
        font_meta = Font(name="Calibri", size=9, italic=True, color="475569")
        font_header = Font(name="Calibri", size=10, bold=True, color=ExcelReportBuilder.COLOR_TEXT_HEADER)
        font_data = Font(name="Calibri", size=10, color="0F172A")
        font_total = Font(name="Calibri", size=10, bold=True, color="0F172A")

        fill_header = PatternFill(start_color=ExcelReportBuilder.COLOR_BG_HEADER, end_color=ExcelReportBuilder.COLOR_BG_HEADER, fill_type="solid")
        fill_alt = PatternFill(start_color=ExcelReportBuilder.COLOR_ROW_ALT, end_color=ExcelReportBuilder.COLOR_ROW_ALT, fill_type="solid")
        fill_total = PatternFill(start_color=ExcelReportBuilder.COLOR_TOTAL, end_color=ExcelReportBuilder.COLOR_TOTAL, fill_type="solid")

        thin_side = Side(border_style="thin", color="CBD5E1")
        border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(top=thin_side, bottom=Side(border_style="double", color="1E40AF"), left=thin_side, right=thin_side)

        current_row = 1

        # 1. Título principal
        ws.cell(row=current_row, column=1, value=f"SISTEMA DE GESTIÓN TM - {title.upper()}").font = font_title
        current_row += 1

        # 2. Metadatos del informe
        meta_items = [f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"]
        if metadata:
            for k, v in metadata.items():
                if v:
                    meta_items.append(f"{k}: {v}")
        meta_str = " | ".join(meta_items)
        ws.cell(row=current_row, column=1, value=meta_str).font = font_meta
        current_row += 2

        # 3. Resumen / Tarjetas si se proveen
        if summary:
            if isinstance(summary, dict):
                summary = [{'label': k, 'value': v} for k, v in summary.items()]
            for item in summary:
                lbl = ws.cell(row=current_row, column=1, value=f"{item.get('label')}:")
                lbl.font = Font(name="Calibri", size=10, bold=True, color="1E40AF")
                val = ws.cell(row=current_row, column=2, value=item.get('value'))
                val.font = Font(name="Calibri", size=10, bold=True)
                current_row += 1
            current_row += 1

        # 4. Encabezados de tabla
        header_row_idx = current_row
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=header_row_idx, column=col_idx, value=str(h_text))
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_cell
        ws.row_dimensions[header_row_idx].height = 24
        current_row += 1

        # 5. Filas de datos
        for r_idx, row_data in enumerate(rows, start=current_row):
            is_alt = (r_idx % 2 == 0)
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.font = font_data
                cell.border = border_cell
                if is_alt:
                    cell.fill = fill_alt

                # Formateo condicional según tipo de dato
                if isinstance(val, (int, float)):
                    # Si parece monetario o porcentaje o cantidad
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, float):
                        cell.number_format = '#,##0.00'
                    else:
                        cell.number_format = '#,##0'
                elif isinstance(val, datetime):
                    cell.value = val.strftime('%Y-%m-%d %H:%M')
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += len(rows)

        # 6. Autoajuste de anchos de columnas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            # Solo evaluar desde la fila de encabezados
            for cell in col:
                if cell.row >= header_row_idx and cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Generar buffer en memoria
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()


class CsvReportBuilder:
    """
    Generador de CSV con BOM UTF-8 para garantizar acentos y caracteres especiales en Excel.
    """
    @staticmethod
    def build_csv(title, headers, rows):
        output = io.StringIO()
        # Escribir BOM UTF-8
        output.write('\ufeff')
        writer = csv.writer(output, dialect='excel')

        # Título y encabezados
        writer.writerow([f"SISTEMA DE GESTIÓN TM - {title}"])
        writer.writerow([f"Generado el: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
        writer.writerow([])
        writer.writerow(headers)

        for row in rows:
            writer.writerow(row)

        return output.getvalue().encode('utf-8-sig')


def create_export_response(format_type, filename_base, title, headers, rows, summary=None, metadata=None):
    """
    Genera la respuesta HTTP adecuada (Excel o CSV) para DRF/Django.
    """
    timestamp = datetime.now().strftime('%Y%m%d_%H%M')
    if format_type == 'excel':
        content = ExcelReportBuilder.build_excel(title, headers, rows, summary, metadata)
        response = HttpResponse(
            content,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename_base}_{timestamp}.xlsx"'
        return response
    elif format_type == 'csv':
        content = CsvReportBuilder.build_csv(title, headers, rows)
        response = HttpResponse(content, content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename_base}_{timestamp}.csv"'
        return response
    return None
